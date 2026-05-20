"""Export merged proposal as PDF (reportlab) or Excel (openpyxl)."""
import io
from typing import Dict, Any, List

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


DEFAULT_SLIDING_SCALE = [
    {"limit": 10_000_000, "pct": 8.0},
    {"limit": 30_000_000, "pct": 6.5},
    {"limit": 80_000_000, "pct": 5.0},
    {"limit": 999_999_999, "pct": 3.5},
]


def _sliding_scale_fee(cc: float, slabs: List[Dict[str, float]]) -> float:
    remaining = max(0.0, float(cc or 0))
    prev = 0.0
    total = 0.0
    for slab in slabs:
        limit = float(slab.get("limit", 0))
        pct = float(slab.get("pct", 0))
        take = min(remaining, max(0.0, limit - prev))
        total += take * (pct / 100.0)
        remaining -= take
        prev = limit
        if remaining <= 0:
            break
    return total


def compute_fee_methods(project: Dict[str, Any], settings: Dict[str, Any] = None) -> Dict[str, Any]:
    """Compute the 4 method values + final fee + phase distribution for export.
    Returns {} if project has no fee_builder."""
    fb = (project or {}).get("fee_builder") or {}
    if not fb:
        return {}
    settings = settings or {}
    details = fb.get("details") or {}
    multipliers = fb.get("multipliers") or {}
    methods = fb.get("methods") or {}

    # Construction cost
    cc = float(details.get("gfa", 0) or 0) * float(details.get("cost_per_m2", 0) or 0)

    # Bottom-up (mirror frontend totals)
    edge_total = 0.0
    active_stages = [s for s in (fb.get("stages_pre") or []) + (fb.get("stages_post") or []) if s.get("on")]
    for mbr in fb.get("members") or []:
        if not mbr.get("on"):
            continue
        for st in active_stages:
            h = float((mbr.get("hours") or {}).get(st["id"], 0) or 0)
            edge_total += h * float(mbr.get("cost_rate", 0) or 0)
    sub_raw = 0.0
    mgmt_on_subs = 0.0
    for sub in fb.get("subs") or []:
        if not sub.get("on"):
            continue
        sub_total = 0.0
        for st in active_stages:
            sub_total += float((sub.get("fees") or {}).get(st["id"], 0) or 0)
        sub_raw += sub_total
        if sub.get("mgmt_on"):
            mgmt_on_subs += sub_total * (float(sub.get("mgmt_pct", 0) or 0) / 100.0)
    pre_cont = edge_total + sub_raw + mgmt_on_subs
    cont_pct = float(multipliers.get("contingency_pct", 0) or 0)
    bottom_up = pre_cont * (1 + cont_pct / 100.0)

    # % of CC
    benchmark_pct = float(methods.get("benchmark_pct") or 9.0)
    pct_fee = cc * (benchmark_pct / 100.0)

    # Sliding scale
    slabs = settings.get("fee_sliding_scale") or DEFAULT_SLIDING_SCALE
    slab_fee = _sliding_scale_fee(cc, slabs)

    # Adjusted
    factors = methods.get("factors") or settings.get("fee_complexity_factors") or []
    bonus = sum(float(f.get("pct", 0) or 0) for f in factors if f.get("on"))
    adjusted = pct_fee * (1 + bonus / 100.0)

    methods_map = {
        "bottom_up": bottom_up,
        "pct_of_cc": pct_fee,
        "sliding":   slab_fee,
        "adjusted":  adjusted,
    }
    method = methods.get("recommended_method") or "bottom_up"
    override = methods.get("final_fee_override")
    if override not in (None, ""):
        try:
            final_fee = float(override)
            final_label = "Manual override"
        except Exception:
            final_fee = methods_map.get(method, bottom_up)
            final_label = method
    else:
        final_fee = methods_map.get(method, bottom_up)
        final_label = {
            "bottom_up": "Bottom-up build-up",
            "pct_of_cc": f"% of Construction Cost ({benchmark_pct:.2f}%)",
            "sliding":   "Sliding-scale slabs",
            "adjusted":  f"Benchmark + {bonus:.0f}% complexity",
        }.get(method, method)

    # Phase distribution: pull from methods.phase_distribution_custom or apply preset
    STAGE_PRESETS = {
        "traditional": {"p1":2,"p2":3,"p3":10,"p4":15,"p5":30,"p6":5,"p7":5,"p8":5,"po1":22,"po2":2,"po3":1},
        "bim_led":     {"p1":2,"p2":4,"p3":14,"p4":18,"p5":25,"p6":4,"p7":3,"p8":4,"po1":22,"po2":3,"po3":1},
        "aia":         {"p1":1,"p2":2,"p3":7, "p4":15,"p5":20,"p6":7,"p7":5,"p8":13,"po1":25,"po2":4,"po3":1},
        "riba":        {"p1":2,"p2":3,"p3":15,"p4":15,"p5":35,"p6":5,"p7":5,"p8":0, "po1":18,"po2":1,"po3":1},
    }
    preset_id = methods.get("phase_preset") or settings.get("fee_phase_preset") or "traditional"
    preset_pct = STAGE_PRESETS.get(preset_id, STAGE_PRESETS["traditional"])
    custom = methods.get("phase_distribution_custom")
    if custom:
        phase_dist = {s["id"]: float(custom.get(s["id"], 0) or 0) for s in active_stages}
    else:
        out = {}
        unmapped = []
        mapped_sum = 0.0
        for s in active_stages:
            if s["id"] in preset_pct:
                out[s["id"]] = float(preset_pct[s["id"]])
                mapped_sum += out[s["id"]]
            else:
                unmapped.append(s["id"])
        leftover = max(0.0, 100.0 - mapped_sum)
        if unmapped:
            share = leftover / len(unmapped)
            for sid in unmapped:
                out[sid] = round(share, 2)
        phase_dist = out

    phase_rows = [
        {"id": s["id"], "name": s["name"], "pct": phase_dist.get(s["id"], 0.0),
         "fee": final_fee * phase_dist.get(s["id"], 0.0) / 100.0}
        for s in active_stages
    ]

    return {
        "currency": details.get("currency") or "USD",
        "construction_cost": cc,
        "methods": [
            {"key": "bottom_up", "label": "A · Bottom-up build-up",   "value": bottom_up},
            {"key": "pct_of_cc", "label": "B · % of Construction Cost", "value": pct_fee},
            {"key": "sliding",   "label": "C · Sliding-scale slabs",    "value": slab_fee},
            {"key": "adjusted",  "label": "D · Benchmark + complexity", "value": adjusted},
        ],
        "selected_method": method,
        "final_fee": final_fee,
        "final_label": final_label,
        "phase_preset": preset_id,
        "phase_distribution": phase_rows,
    }


def _style():
    s = getSampleStyleSheet()
    s.add(ParagraphStyle(name="H", fontName="Helvetica-Bold", fontSize=22, leading=26, spaceAfter=8))
    s.add(ParagraphStyle(name="Sub", fontName="Helvetica", fontSize=10, leading=12, textColor=colors.grey, spaceAfter=14))
    s.add(ParagraphStyle(name="Section", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=14, spaceAfter=6, textColor=colors.HexColor("#0055FF")))
    s.add(ParagraphStyle(name="Body", fontName="Helvetica", fontSize=9, leading=12))
    return s


def _table(data, col_widths=None, header=True):
    t = Table(data, colWidths=col_widths)
    style = [
        ("FONTNAME", (0,0), (-1,-1), "Helvetica"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LINEBELOW", (0,0), (-1,-1), 0.25, colors.HexColor("#E4E4E7")),
        ("LEFTPADDING", (0,0), (-1,-1), 4),
        ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]
    if header:
        style += [
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0A0A0B")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ]
    t.setStyle(TableStyle(style))
    return t


def export_pdf(project: Dict[str, Any], invites: List[Dict[str, Any]], settings: Dict[str, Any] = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
    s = _style()
    flow = []
    flow.append(Paragraph(project.get("title", "Untitled RFP"), s["H"]))
    flow.append(Paragraph(f"Client: {project.get('client_name') or '—'} · Status: {project.get('status', '—')}", s["Sub"]))

    a = project.get("analysis") or {}
    if a.get("summary"):
        flow.append(Paragraph("Executive Summary", s["Section"]))
        flow.append(Paragraph(a["summary"], s["Body"]))

    # Submitted fees
    submitted = [i for i in invites if i.get("status") == "submitted" and i.get("fee") is not None]
    flow.append(Paragraph("Merged Fee Proposal", s["Section"]))
    if submitted:
        rows = [["Discipline", "Consultant", "Fee", "Currency", "Timeline"]]
        for i in submitted:
            rows.append([
                i.get("discipline", ""),
                f"{i.get('consultant_name','')}\n{i.get('consultant_company','')}",
                f"{float(i['fee']):,.2f}",
                i.get("currency", ""),
                i.get("response_timeline", ""),
            ])
        flow.append(_table(rows, col_widths=[35*mm, 50*mm, 25*mm, 20*mm, 30*mm]))
        # Totals
        totals: Dict[str, float] = {}
        for i in submitted:
            cur = i.get("currency") or "USD"
            totals[cur] = totals.get(cur, 0.0) + float(i["fee"])
        flow.append(Spacer(1, 8))
        for cur, tot in totals.items():
            flow.append(Paragraph(f"<b>Total {cur}: {tot:,.2f}</b>", s["Body"]))
    else:
        flow.append(Paragraph("<i>No fees submitted yet.</i>", s["Body"]))

    # Fee Methods comparison + final fee (only if Fee Builder has been used)
    fm = compute_fee_methods(project, settings)
    if fm:
        flow.append(Paragraph("Fee Methods · Comparison &amp; Rationale", s["Section"]))
        cur = fm["currency"]
        flow.append(Paragraph(f"Construction cost: <b>{cur} {fm['construction_cost']:,.0f}</b>", s["Body"]))
        rows = [["Method", "Fee", "vs Bottom-up"]]
        bu = next((m["value"] for m in fm["methods"] if m["key"] == "bottom_up"), 0)
        for mm_ in fm["methods"]:
            delta = ((mm_["value"] - bu) / bu * 100.0) if bu > 0 else 0.0
            d_str = "—" if mm_["key"] == "bottom_up" else f"{'+' if delta >= 0 else ''}{delta:.1f}%"
            rows.append([mm_["label"], f"{cur} {mm_['value']:,.0f}", d_str])
        flow.append(_table(rows, col_widths=[80*mm, 50*mm, 30*mm]))
        flow.append(Spacer(1, 8))
        flow.append(Paragraph(
            f"<b>Final fee: {cur} {fm['final_fee']:,.0f}</b> &nbsp;·&nbsp; <font color='#666666'>{fm['final_label']}</font>",
            s["Body"]
        ))
        if fm["phase_distribution"]:
            flow.append(Spacer(1, 8))
            flow.append(Paragraph(f"Phase distribution · preset: {fm['phase_preset']}", s["Body"]))
            ph_rows = [["Stage", "% of fee", "Fee"]]
            for r in fm["phase_distribution"]:
                ph_rows.append([r["name"], f"{r['pct']:.2f}%", f"{cur} {r['fee']:,.0f}"])
            flow.append(_table(ph_rows, col_widths=[80*mm, 40*mm, 40*mm]))

    if a.get("key_dates"):
        flow.append(PageBreak())
        flow.append(Paragraph("Key Dates", s["Section"]))
        rows = [["Date", "Label", "Type"]]
        for d in a["key_dates"]:
            rows.append([d.get("date",""), d.get("label",""), d.get("type","")])
        flow.append(_table(rows, col_widths=[30*mm, 90*mm, 30*mm]))

    if a.get("requirements"):
        flow.append(Paragraph("Requirements", s["Section"]))
        rows = [["ID", "Category", "Requirement", "Mandatory"]]
        for r in a["requirements"]:
            rows.append([r.get("id",""), r.get("category",""),
                         Paragraph(r.get("requirement",""), s["Body"]),
                         "Yes" if r.get("mandatory") else "No"])
        flow.append(_table(rows, col_widths=[18*mm, 30*mm, 105*mm, 18*mm]))

    if a.get("disciplines"):
        flow.append(Paragraph("Disciplines", s["Section"]))
        rows = [["Name", "Scope"]]
        for d in a["disciplines"]:
            rows.append([d.get("name",""), Paragraph(d.get("scope_summary",""), s["Body"])])
        flow.append(_table(rows, col_widths=[40*mm, 130*mm]))

    doc.build(flow)
    return buf.getvalue()


def export_xlsx(project: Dict[str, Any], invites: List[Dict[str, Any]], settings: Dict[str, Any] = None) -> bytes:
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="0A0A0B")
    head_font = Font(bold=True, color="FFFFFF", name="Calibri")
    border = Border(*[Side(style="thin", color="E4E4E7")]*4)

    def write_table(ws, headers, rows, widths=None):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.fill = head_fill; c.font = head_font; c.alignment = Alignment(vertical="center")
            if widths and ci-1 < len(widths): ws.column_dimensions[c.column_letter].width = widths[ci-1]
        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.border = border; c.alignment = Alignment(vertical="top", wrap_text=True)

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"
    a = project.get("analysis") or {}
    ws["A1"] = "RFP Analyser — Merged Proposal"; ws["A1"].font = Font(bold=True, size=16)
    ws["A3"] = "Title"; ws["B3"] = project.get("title", "")
    ws["A4"] = "Client"; ws["B4"] = project.get("client_name", "")
    ws["A5"] = "Status"; ws["B5"] = project.get("status", "")
    ws["A6"] = "Summary"; ws["B6"] = a.get("summary", "")
    for col in ["A", "B"]: ws.column_dimensions[col].width = 30 if col == "A" else 80

    # Sheet 2: Fees
    ws2 = wb.create_sheet("Fees")
    submitted = [i for i in invites if i.get("status") == "submitted" and i.get("fee") is not None]
    rows = [[i.get("discipline",""), i.get("consultant_name",""), i.get("consultant_company",""),
             float(i["fee"]), i.get("currency",""), i.get("response_timeline",""), i.get("response_notes","")]
            for i in submitted]
    write_table(ws2, ["Discipline","Consultant","Company","Fee","Currency","Timeline","Notes"], rows, [22, 22, 22, 14, 10, 18, 40])

    # Totals
    totals: Dict[str, float] = {}
    for i in submitted:
        cur = i.get("currency") or "USD"
        totals[cur] = totals.get(cur, 0.0) + float(i["fee"])
    last = len(rows) + 3
    ws2.cell(row=last, column=1, value="TOTALS").font = Font(bold=True)
    for j, (cur, tot) in enumerate(totals.items(), 1):
        ws2.cell(row=last+j, column=1, value=cur).font = Font(bold=True)
        ws2.cell(row=last+j, column=4, value=tot).font = Font(bold=True)

    # Sheet 3: Requirements
    ws3 = wb.create_sheet("Requirements")
    rows = [[r.get("id",""), r.get("category",""), r.get("requirement",""),
             "Yes" if r.get("mandatory") else "No", r.get("source","")]
            for r in (a.get("requirements") or [])]
    write_table(ws3, ["ID","Category","Requirement","Mandatory","Source"], rows, [10, 22, 60, 12, 22])

    # Sheet 4: Key dates
    ws4 = wb.create_sheet("Key Dates")
    rows = [[d.get("date",""), d.get("label",""), d.get("type","")] for d in (a.get("key_dates") or [])]
    write_table(ws4, ["Date","Label","Type"], rows, [16, 60, 16])

    # Sheet 5: Disciplines
    ws5 = wb.create_sheet("Disciplines")
    rows = [[d.get("name",""), d.get("description",""), d.get("scope_summary","")] for d in (a.get("disciplines") or [])]
    write_table(ws5, ["Name","Description","Scope summary"], rows, [22, 40, 60])

    # Sheet 6: Fee Methods (only if fee_builder filled in)
    fm = compute_fee_methods(project, settings)
    if fm:
        ws6 = wb.create_sheet("Fee Methods")
        cur = fm["currency"]
        ws6["A1"] = "Fee Methods · Comparison"
        ws6["A1"].font = Font(bold=True, size=14)
        ws6["A3"] = "Construction cost"; ws6["B3"] = fm["construction_cost"]
        ws6["A4"] = "Selected method";    ws6["B4"] = fm["final_label"]
        ws6["A5"] = "Final fee";          ws6["B5"] = fm["final_fee"]
        ws6["A6"] = "Currency";           ws6["B6"] = cur
        # comparison
        ws6["A8"] = "Method"; ws6["B8"] = "Fee"; ws6["C8"] = "vs Bottom-up"
        for col in ("A8", "B8", "C8"):
            ws6[col].fill = head_fill; ws6[col].font = head_font
        bu = next((m["value"] for m in fm["methods"] if m["key"] == "bottom_up"), 0)
        for idx, m_ in enumerate(fm["methods"], start=9):
            ws6.cell(row=idx, column=1, value=m_["label"])
            ws6.cell(row=idx, column=2, value=m_["value"])
            delta = ((m_["value"] - bu) / bu * 100.0) if bu > 0 else 0.0
            ws6.cell(row=idx, column=3, value=("—" if m_["key"] == "bottom_up" else f"{delta:+.1f}%"))
        # phase distribution
        ws6.cell(row=15, column=1, value=f"Phase distribution · {fm['phase_preset']}").font = Font(bold=True)
        ws6["A17"] = "Stage"; ws6["B17"] = "% of fee"; ws6["C17"] = "Fee"
        for col in ("A17", "B17", "C17"):
            ws6[col].fill = head_fill; ws6[col].font = head_font
        for i, r in enumerate(fm["phase_distribution"], start=18):
            ws6.cell(row=i, column=1, value=r["name"])
            ws6.cell(row=i, column=2, value=r["pct"] / 100.0).number_format = "0.00%"
            ws6.cell(row=i, column=3, value=r["fee"])
        ws6.column_dimensions["A"].width = 50
        ws6.column_dimensions["B"].width = 18
        ws6.column_dimensions["C"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
