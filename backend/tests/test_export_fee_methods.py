"""Tests for compute_fee_methods + PDF/XLSX export with fee_builder.

Covers iteration_5 review request:
- compute_fee_methods math (4 methods + final + sliding scale + override)
- Phase distribution honoring active stages and preset/custom override
- /api/projects/{id}/export?format=pdf and =xlsx for projects WITH and WITHOUT fee_builder
"""
import io
import os
import uuid
import zipfile
import requests
import pytest
from openpyxl import load_workbook

import sys
sys.path.insert(0, "/app/backend")
from export_service import compute_fee_methods  # noqa: E402

def _load_backend_url():
    v = os.environ.get("REACT_APP_BACKEND_URL")
    if v:
        return v.rstrip("/")
    try:
        with open("/app/frontend/.env") as f:
            for line in f:
                if line.startswith("REACT_APP_BACKEND_URL="):
                    return line.split("=", 1)[1].strip().rstrip("/")
    except FileNotFoundError:
        pass
    return ""

BASE_URL = _load_backend_url()


# ------------------- Unit tests: compute_fee_methods ------------------- #

def _mk_project(fee_builder=None, analysis=None):
    return {
        "id": "p1",
        "title": "Test",
        "client_name": "ClientCo",
        "status": "draft",
        "analysis": analysis or {},
        "fee_builder": fee_builder,
    }


def _seeded_fb(extra_methods=None):
    fb = {
        "details": {"gfa": 5000, "cost_per_m2": 10000, "currency": "AED"},
        "multipliers": {"contingency_pct": 0, "vat_pct": 0},
        "stages_pre": [
            {"id": "p1", "name": "Mobilization", "on": False},
            {"id": "p3", "name": "Concept Design", "on": True},
            {"id": "p5", "name": "Detailed Design", "on": True},
        ],
        "stages_post": [
            {"id": "po1", "name": "Construction Supervision", "on": True},
        ],
        "members": [],
        "subs": [],
        "methods": {
            "recommended_method": "adjusted",
            "phase_preset": "traditional",
            "benchmark_pct": 12,
            "factors": [{"name": "Healthcare", "pct": 40, "on": True}],
        },
    }
    if extra_methods:
        fb["methods"].update(extra_methods)
    return fb


class TestComputeFeeMethods:
    def test_empty_project_returns_empty(self):
        assert compute_fee_methods(_mk_project(None)) == {}
        assert compute_fee_methods(_mk_project({})) == {}

    def test_method_d_adjusted(self):
        out = compute_fee_methods(_mk_project(_seeded_fb()), {})
        # Method D = 50M × 0.12 × 1.40 = 8,400,000
        d = next(m for m in out["methods"] if m["key"] == "adjusted")
        assert d["value"] == pytest.approx(8_400_000, rel=1e-6)
        # Method B = 50M × 0.12 = 6,000,000
        b = next(m for m in out["methods"] if m["key"] == "pct_of_cc")
        assert b["value"] == pytest.approx(6_000_000, rel=1e-6)

    def test_method_c_sliding_default_slabs(self):
        out = compute_fee_methods(_mk_project(_seeded_fb()), {})
        # CC=50M -> 10M*8% + 20M*6.5% + 20M*5% = 800k + 1.3M + 1M = 3,100,000
        c = next(m for m in out["methods"] if m["key"] == "sliding")
        assert c["value"] == pytest.approx(3_100_000, rel=1e-6)

    def test_final_fee_equals_selected_method(self):
        out = compute_fee_methods(_mk_project(_seeded_fb()), {})
        assert out["final_fee"] == pytest.approx(8_400_000, rel=1e-6)
        assert out["selected_method"] == "adjusted"
        assert "complexity" in out["final_label"].lower()

    def test_final_fee_override_wins(self):
        fb = _seeded_fb({"final_fee_override": 7_000_000})
        out = compute_fee_methods(_mk_project(fb), {})
        assert out["final_fee"] == pytest.approx(7_000_000, rel=1e-6)
        assert out["final_label"] == "Manual override"

    def test_phase_distribution_active_stages_preset(self):
        # Only p3, p5, po1 active; traditional preset => 10/30/22
        out = compute_fee_methods(_mk_project(_seeded_fb()), {})
        pd = {r["id"]: r["pct"] for r in out["phase_distribution"]}
        assert set(pd.keys()) == {"p3", "p5", "po1"}
        assert pd["p3"] == 10
        assert pd["p5"] == 30
        assert pd["po1"] == 22

    def test_phase_distribution_custom_override(self):
        fb = _seeded_fb({"phase_distribution_custom": {"p3": 20, "p5": 50, "po1": 30}})
        out = compute_fee_methods(_mk_project(fb), {})
        pd = {r["id"]: r["pct"] for r in out["phase_distribution"]}
        assert pd == {"p3": 20.0, "p5": 50.0, "po1": 30.0}
        # Phase fee calculations track the override
        po1_row = next(r for r in out["phase_distribution"] if r["id"] == "po1")
        assert po1_row["fee"] == pytest.approx(8_400_000 * 0.30, rel=1e-6)


# ------------------- Integration: export endpoints ------------------- #

@pytest.fixture(scope="module")
def auth_session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    r = s.post(f"{BASE_URL}/api/auth/login", json={
        "email": "demo@rfpqa.example.com",
        "password": "DemoPass123!",
    })
    if r.status_code != 200:
        pytest.skip(f"demo login failed: {r.status_code} {r.text[:120]}")
    s.headers["Authorization"] = f"Bearer {r.json()['token']}"
    return s


@pytest.fixture(scope="module")
def project_without_fb(auth_session):
    """Create a fresh project with NO fee_builder to verify back-compat exports."""
    r = auth_session.post(f"{BASE_URL}/api/projects", json={
        "title": f"TEST_export_nofb_{uuid.uuid4().hex[:6]}",
        "client_name": "TEST_Client",
    })
    assert r.status_code in (200, 201), r.text
    pid = r.json()["id"]
    yield pid
    auth_session.delete(f"{BASE_URL}/api/projects/{pid}")


@pytest.fixture(scope="module")
def project_with_fb(auth_session):
    """Create a fresh project then seed its fee_builder via PATCH."""
    r = auth_session.post(f"{BASE_URL}/api/projects", json={
        "title": f"TEST_export_fb_{uuid.uuid4().hex[:6]}",
        "client_name": "TEST_Client",
    })
    assert r.status_code in (200, 201), r.text
    pid = r.json()["id"]
    # Seed fee_builder
    fb = _seeded_fb()
    pr = auth_session.patch(f"{BASE_URL}/api/projects/{pid}/fee-builder", json=fb)
    assert pr.status_code in (200, 204), pr.text
    yield pid
    auth_session.delete(f"{BASE_URL}/api/projects/{pid}")


class TestExportBackCompat:
    def test_pdf_export_without_fee_builder(self, auth_session, project_without_fb):
        r = auth_session.get(f"{BASE_URL}/api/projects/{project_without_fb}/export",
                             params={"format": "pdf"})
        assert r.status_code == 200, r.text[:200]
        assert r.content.startswith(b"%PDF-"), "not a PDF"
        assert len(r.content) > 1500, f"PDF too small: {len(r.content)} bytes"

    def test_xlsx_export_without_fee_builder(self, auth_session, project_without_fb):
        r = auth_session.get(f"{BASE_URL}/api/projects/{project_without_fb}/export",
                             params={"format": "xlsx"})
        assert r.status_code == 200, r.text[:200]
        assert r.content[:2] == b"PK", "not a zip/xlsx"
        wb = load_workbook(io.BytesIO(r.content))
        # No "Fee Methods" sheet when fee_builder absent
        assert "Fee Methods" not in wb.sheetnames, f"unexpected sheet: {wb.sheetnames}"


class TestExportWithFeeBuilder:
    def test_pdf_export_with_fee_builder(self, auth_session, project_with_fb):
        r = auth_session.get(f"{BASE_URL}/api/projects/{project_with_fb}/export",
                             params={"format": "pdf"})
        assert r.status_code == 200, r.text[:200]
        assert r.content.startswith(b"%PDF-")
        assert len(r.content) > 2048

    def test_xlsx_export_fee_methods_sheet_values(self, auth_session, project_with_fb):
        r = auth_session.get(f"{BASE_URL}/api/projects/{project_with_fb}/export",
                             params={"format": "xlsx"})
        assert r.status_code == 200
        assert r.content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(r.content))
        assert "Fee Methods" in wb.sheetnames
        ws = wb["Fee Methods"]
        # B3 construction cost = 50_000_000
        assert ws["B3"].value == pytest.approx(50_000_000)
        # B5 final fee = 8,400,000
        assert ws["B5"].value == pytest.approx(8_400_000, rel=1e-6)
        # Selected method label includes complexity
        assert "complexity" in str(ws["B4"].value).lower()
        # Three phase rows starting at row 18 — p3 Concept Design = 10%
        stage_names = [ws.cell(row=18+i, column=1).value for i in range(3)]
        pct_values = [ws.cell(row=18+i, column=2).value for i in range(3)]
        assert "Concept Design" in stage_names
        assert "Detailed Design" in stage_names
        assert "Construction Supervision" in stage_names
        # XLSX stores as 0.xxx with 0.00% number format
        assert any(abs((v or 0) - 0.30) < 1e-6 for v in pct_values), pct_values

    def test_xlsx_export_custom_phase_distribution(self, auth_session, project_with_fb):
        # Patch in custom phase distribution
        custom = {"p3": 20, "p5": 50, "po1": 30}
        pr = auth_session.patch(f"{BASE_URL}/api/projects/{project_with_fb}/fee-builder",
                                json=_seeded_fb({"phase_distribution_custom": custom}))
        assert pr.status_code in (200, 204)
        r = auth_session.get(f"{BASE_URL}/api/projects/{project_with_fb}/export",
                             params={"format": "xlsx"})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["Fee Methods"]
        pct_values = [ws.cell(row=18+i, column=2).value for i in range(3)]
        # Custom values 20/50/30 => 0.20, 0.50, 0.30
        assert any(abs((v or 0) - 0.50) < 1e-6 for v in pct_values), pct_values
        assert any(abs((v or 0) - 0.20) < 1e-6 for v in pct_values), pct_values

    def test_xlsx_export_override_final_fee(self, auth_session, project_with_fb):
        pr = auth_session.patch(f"{BASE_URL}/api/projects/{project_with_fb}/fee-builder",
                                json=_seeded_fb({"final_fee_override": 7_000_000}))
        assert pr.status_code in (200, 204)
        r = auth_session.get(f"{BASE_URL}/api/projects/{project_with_fb}/export",
                             params={"format": "xlsx"})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["Fee Methods"]
        assert ws["B5"].value == pytest.approx(7_000_000)
        assert "override" in str(ws["B4"].value).lower()
